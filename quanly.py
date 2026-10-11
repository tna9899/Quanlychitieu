import os
import io
import sys
import json
import ctypes
import threading
import webbrowser  # Thư viện để mở link trình duyệt
import tkinter as tk
from tkinter import messagebox, ttk
from datetime import datetime
import pandas as pd
from tkcalendar import DateEntry
from PIL import Image, ImageDraw, ImageTk

# Thiết lập AppUserModelID cho Windows:
# CHỈ gọi khi chạy bằng python (dev mode).
# Khi đã đóng gói thành file .exe (sys.frozen = True), KHÔNG gọi lệnh này để tránh làm mất icon nhúng sẵn của file .exe trên Taskbar.
if sys.platform.startswith('win'):
    if not getattr(sys, 'frozen', False):
        try:
            app_id = 'trinhngocanh.quanlychitieu.app.1.0'
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(app_id)
        except Exception:
            pass

# Thư viện định dạng Excel trước khi upload
import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
# Thư viện vẽ biểu đồ của openpyxl cho Google Sheets
from openpyxl.chart import PieChart, Reference

# Các thư viện kết nối Google Drive API
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaInMemoryUpload

# Phạm vi quyền truy cập
SCOPES = ['https://www.googleapis.com/auth/drive.file']
# Tên file cơ sở lưu trên Google Drive
DRIVE_FILE_BASE_NAME = "Danh sách chi tiêu"
DEFAULT_DRIVE_FILE_ID = "1kMZ7lLpjMlmcOkRBZhFgby4TnTmwgYn2jbaK1z3wW00"


def get_drive_file_name_for_year(year: int) -> str:
    """Tạo tên file theo từng năm, ví dụ: 'Danh sách chi tiêu 2026'"""
    return f"{DRIVE_FILE_BASE_NAME} {year}"


def get_categories_for_month(month: int, year: int = 2026):
    """
    Quy tắc danh mục:
    - Từ tháng 10-2026 trở đi: Chuyển 'Điện nước' thành 'Bố Mẹ'
    - Các tháng trước tháng 10: Giữ nguyên 'Điện nước'
    """
    if year > 2026 or (year == 2026 and month >= 10):
        return [
            "Ăn uống", "Di chuyển", "Mua sắm", "Giải trí",
            "Bố Mẹ", "Cá nhân", "Tiền trọ", "Hiếu/ Hỉ"
        ]
    return [
        "Ăn uống", "Di chuyển", "Mua sắm", "Giải trí",
        "Điện nước", "Cá nhân", "Tiền trọ", "Hiếu/ Hỉ"
    ]


# Dữ liệu cấu hình mẫu (người dùng sẽ cung cấp qua credentials.json và token.json để bảo mật)
DEFAULT_CREDENTIALS_INFO = {
    "installed": {
        "client_id": "YOUR_CLIENT_ID.apps.googleusercontent.com",
        "project_id": "quanlychitieu",
        "auth_uri": "https://accounts.google.com/o/oauth2/auth",
        "token_uri": "https://oauth2.googleapis.com/token",
        "auth_provider_x509_cert_url": "https://www.googleapis.com/oauth2/v1/certs",
        "client_secret": "YOUR_CLIENT_SECRET",
        "redirect_uris": ["http://localhost"]
    }
}

DEFAULT_TOKEN_INFO = {
    "token": "",
    "refresh_token": "",
    "token_uri": "https://oauth2.googleapis.com/token",
    "client_id": "",
    "client_secret": "",
    "scopes": ["https://www.googleapis.com/auth/drive.file"],
    "universe_domain": "googleapis.com",
    "account": "",
    "expiry": ""
}


def get_app_data_dir():
    """Lấy thư mục AppData chuyên dụng để lưu token/cache, tuyệt đối không làm rác thư mục chứa file .exe"""
    app_data = os.getenv('APPDATA')
    if app_data:
        path = os.path.join(app_data, 'QuanLyChiTieu')
    else:
        path = os.path.join(os.path.expanduser('~'), '.quanlychitieu')
    os.makedirs(path, exist_ok=True)
    return path


def get_resource_path(relative_path):
    """Lấy đường dẫn tuyệt đối cho tài nguyên (dùng được cả khi chạy dev và khi đóng gói .exe)"""
    try:
        # Khi đóng gói bằng PyInstaller với --onefile
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.dirname(os.path.abspath(__file__))
    
    # 1. Thử tìm trong gói nhúng tạm thời _MEIPASS
    full_path = os.path.join(base_path, relative_path)
    if os.path.exists(full_path):
        return full_path
    
    # 2. Thử tìm trong AppData
    appdata_path = os.path.join(get_app_data_dir(), relative_path)
    if os.path.exists(appdata_path):
        return appdata_path

    # 3. Nếu không có, tìm ở thư mục cùng cấp với file .exe (hoặc dev script)
    exe_dir = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
    return os.path.join(exe_dir, relative_path)


def make_rounded_button_images(width, height, radius, fill_normal, fill_hover, outline_normal=None, outline_hover=None):
    """Tạo ảnh nền bo tròn góc mịn (khử răng cưa 2x) cho nút bấm"""
    scale = 2
    w_s, h_s, r_s = width * scale, height * scale, radius * scale
    
    # Normal state
    img_normal = Image.new('RGBA', (w_s, h_s), (0, 0, 0, 0))
    d_normal = ImageDraw.Draw(img_normal)
    d_normal.rounded_rectangle(
        (0, 0, w_s - 1, h_s - 1),
        radius=r_s,
        fill=fill_normal,
        outline=outline_normal,
        width=scale if outline_normal else 0
    )
    img_normal = img_normal.resize((width, height), Image.Resampling.LANCZOS)
    photo_normal = ImageTk.PhotoImage(img_normal)
    
    # Hover state
    img_hover = Image.new('RGBA', (w_s, h_s), (0, 0, 0, 0))
    d_hover = ImageDraw.Draw(img_hover)
    d_hover.rounded_rectangle(
        (0, 0, w_s - 1, h_s - 1),
        radius=r_s,
        fill=fill_hover,
        outline=outline_hover or outline_normal,
        width=scale if (outline_hover or outline_normal) else 0
    )
    img_hover = img_hover.resize((width, height), Image.Resampling.LANCZOS)
    photo_hover = ImageTk.PhotoImage(img_hover)
    
    return photo_normal, photo_hover


class ExpenditureApp:

    def __init__(self, root):
        self.root = root
        self.root.title("Quản lý Chi tiêu")
        self.root.configure(bg="#F8FAFC")
        
        # --- CẤU HÌNH ICON CHO CỬA SỔ APP & TASKBAR ---
        icon_path = get_resource_path("logo.ico")
        png_path = get_resource_path("logo.png")
        self.icon_photo = None
        self.header_logo = None

        # 1. Thử thiết lập icon bằng iconbitmap (chuẩn Windows Win32)
        if os.path.exists(icon_path):
            try:
                self.root.iconbitmap(default=icon_path)
            except Exception:
                pass
            try:
                self.root.iconbitmap(icon_path)
            except Exception:
                pass

        # 2. Thiết lập icon bằng iconphoto (hỗ trợ Tkinter RGBA trên Taskbar/Alt-Tab)
        if os.path.exists(png_path):
            try:
                self.icon_photo = tk.PhotoImage(file=png_path)
                self.root.iconphoto(True, self.icon_photo)
            except Exception:
                pass
        elif os.path.exists(icon_path):
            try:
                self.icon_photo = ImageTk.PhotoImage(file=icon_path)
                self.root.iconphoto(True, self.icon_photo)
            except Exception:
                pass

        # 3. Ép cập nhật icon trực tiếp lên thanh tiêu đề qua Windows Win32 API
        def apply_win32_titlebar_icon():
            if sys.platform.startswith('win'):
                try:
                    hwnd = ctypes.windll.user32.GetParent(self.root.winfo_id())
                    if not hwnd:
                        hwnd = self.root.winfo_id()
                    WM_SETICON = 0x0080
                    if os.path.exists(icon_path):
                        hicon_small = ctypes.windll.user32.LoadImageW(None, icon_path, 1, 16, 16, 0x0010)
                        hicon_big = ctypes.windll.user32.LoadImageW(None, icon_path, 1, 32, 32, 0x0010)
                        if hicon_small:
                            ctypes.windll.user32.SendMessageW(hwnd, WM_SETICON, 0, hicon_small)
                        if hicon_big:
                            ctypes.windll.user32.SendMessageW(hwnd, WM_SETICON, 1, hicon_big)
                except Exception:
                    pass

        self.root.after(50, apply_win32_titlebar_icon)
        self.root.after(200, apply_win32_titlebar_icon)
        
        # --- CẤU HÌNH KÍCH THƯỚC VÀ CĂN GIỮA MÀN HÌNH ---
        window_width = 390
        window_height = 610
        
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        
        center_x = int((screen_width - window_width) / 2)
        center_y = int((screen_height - window_height) / 2)
        
        self.root.geometry(f"{window_width}x{window_height}+{center_x}+{center_y}")
        self.root.resizable(False, False)

        if getattr(sys, 'frozen', False):
            self.current_dir = os.path.dirname(sys.executable)
        else:
            self.current_dir = os.path.dirname(os.path.abspath(__file__))

        self.current_year = datetime.now().year
        # Danh mục tự động nhận diện theo tháng (Tháng >= 10 có 'Bố Mẹ')
        self.fixed_categories = get_categories_for_month(datetime.now().month, self.current_year)
        
        # DataFrame quản lý các cột chi tiêu cơ bản
        self.df = pd.DataFrame(columns=["Ngày", "Danh mục", "Số tiền", "Ghi chú"])

        # --- 1. KHỞI TẠO BIẾN TKINTER ---
        self.amount_var = tk.StringVar()
        self.amount_var.trace_add("write", self.format_amount_live)

        # Cấu hình theme TTK hiện đại
        style = ttk.Style()
        style.theme_use('clam')
        style.configure(
            "Custom.TCombobox",
            fieldbackground="#FFFFFF",
            background="#F1F5F9",
            foreground="#1E293B",
            padding=5,
            arrowsize=14,
            relief="flat"
        )
        style.map(
            "Custom.TCombobox",
            fieldbackground=[('readonly', '#FFFFFF')],
            selectbackground=[('readonly', '#FFFFFF')],
            selectforeground=[('readonly', '#1E293B')]
        )

        # --- 2. DỰNG GIAO DIỆN HIỆN ĐẠI (MODERN CARD UI) ---
        main_container = tk.Frame(root, bg="#F8FAFC", padx=16, pady=12)
        main_container.pack(fill=tk.BOTH, expand=True)

        # A. KHỐI TIÊU ĐỀ (HEADER CARD)
        header_frame = tk.Frame(
            main_container,
            bg="#FFFFFF",
            highlightthickness=1,
            highlightbackground="#E2E8F0",
            padx=12,
            pady=10
        )
        header_frame.pack(fill=tk.X, pady=(0, 10))

        # Hiển thị Logo thu nhỏ trên Header
        logo_38_path = get_resource_path("logo_38.png")
        if os.path.exists(logo_38_path):
            try:
                self.header_logo = tk.PhotoImage(file=logo_38_path)
                logo_label = tk.Label(header_frame, image=self.header_logo, bg="#FFFFFF")
                logo_label.pack(side=tk.LEFT, padx=(0, 10))
            except Exception:
                pass
        elif os.path.exists(png_path):
            try:
                logo_pil = Image.open(png_path).resize((38, 38), Image.Resampling.LANCZOS)
                self.header_logo = ImageTk.PhotoImage(logo_pil)
                logo_label = tk.Label(header_frame, image=self.header_logo, bg="#FFFFFF")
                logo_label.pack(side=tk.LEFT, padx=(0, 10))
            except Exception:
                pass
        elif os.path.exists(icon_path):
            try:
                logo_pil = Image.open(icon_path).resize((38, 38), Image.Resampling.LANCZOS)
                self.header_logo = ImageTk.PhotoImage(logo_pil)
                logo_label = tk.Label(header_frame, image=self.header_logo, bg="#FFFFFF")
                logo_label.pack(side=tk.LEFT, padx=(0, 10))
            except Exception:
                pass

        title_box = tk.Frame(header_frame, bg="#FFFFFF")
        title_box.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        tk.Label(
            title_box,
            text="QUẢN LÝ CHI TIÊU",
            font=("Segoe UI", 11, "bold"),
            fg="#0F172A",
            bg="#FFFFFF"
        ).pack(anchor=tk.W)

        tk.Label(
            title_box,
            text="Đồng bộ đám mây Google Drive",
            font=("Segoe UI", 8),
            fg="#64748B",
            bg="#FFFFFF"
        ).pack(anchor=tk.W)

        # Thông tin tài khoản đã đồng bộ (in nghiêng)
        self.account_label = tk.Label(
            title_box,
            text="",
            font=("Segoe UI", 8, "italic"),
            fg="#64748B",
            bg="#FFFFFF"
        )
        self.account_label.pack(anchor=tk.W)

        # Huy hiệu trạng thái kết nối
        self.status_badge = tk.Label(
            header_frame,
            text="● Đang nạp...",
            font=("Segoe UI", 8, "bold"),
            fg="#F59E0B",
            bg="#FEF3C7",
            padx=7,
            pady=3
        )
        self.status_badge.pack(side=tk.RIGHT, padx=2)

        # B. THẺ NHẬP LIỆU (FORM CARD)
        card_frame = tk.Frame(
            main_container,
            bg="#FFFFFF",
            highlightthickness=1,
            highlightbackground="#E2E8F0",
            padx=14,
            pady=12
        )
        card_frame.pack(fill=tk.X, pady=(0, 10))

        # 1. Danh mục
        tk.Label(
            card_frame,
            text="🏷️  Danh mục chi tiêu:",
            font=("Segoe UI", 9, "bold"),
            fg="#334155",
            bg="#FFFFFF"
        ).pack(anchor=tk.W, pady=(0, 3))
        
        self.category_cb = ttk.Combobox(
            card_frame,
            values=self.fixed_categories,
            font=("Segoe UI", 9),
            state="readonly",
            style="Custom.TCombobox"
        )
        self.category_cb.pack(fill=tk.X, pady=(0, 10))
        self.category_cb.current(0)

        # Chặn bôi đen khi click chọn hoặc chuyển focus trong danh mục chi tiêu
        def _clear_cb_selection(e=None):
            self.category_cb.after(10, self.category_cb.selection_clear)

        self.category_cb.bind("<<ComboboxSelected>>", _clear_cb_selection)
        self.category_cb.bind("<FocusIn>", _clear_cb_selection)
        self.category_cb.bind("<ButtonRelease-1>", _clear_cb_selection)

        # 2. Ngày giao dịch
        tk.Label(
            card_frame,
            text="📅  Chọn ngày giao dịch:",
            font=("Segoe UI", 9, "bold"),
            fg="#334155",
            bg="#FFFFFF"
        ).pack(anchor=tk.W, pady=(0, 3))

        self.date_entry = DateEntry(
            card_frame,
            font=("Segoe UI", 9),
            background="#2563EB",
            foreground="#FFFFFF",
            headersbackground="#1E40AF",
            headersforeground="#FFFFFF",
            selectbackground="#3B82F6",
            selectforeground="#FFFFFF",
            normalbackground="#FFFFFF",
            normalforeground="#1E293B",
            weekendbackground="#F8FAFC",
            weekendforeground="#DC2626",
            borderwidth=1,
            date_pattern="yyyy-mm-dd",
            state="readonly",
            cursor="hand2"
        )
        self.date_entry.configure(cursor="hand2")
        self.date_entry.pack(fill=tk.X, pady=(0, 10), ipady=3)

        # Làm nổi bật ngày hôm nay trên lịch chọn ngày
        try:
            today_date = datetime.now().date()
            self.date_entry._calendar.calevent_create(today_date, "Hôm nay", "today_tag")
            self.date_entry._calendar.tag_config(
                "today_tag",
                background="#FEF08A",
                foreground="#854D0E"
            )
        except Exception:
            pass

        # Bắt sự kiện click để mở popup lịch giống như dropdown Combobox (không hiện con trỏ văn bản)
        def _open_calendar_popup(event=None):
            if not getattr(self.date_entry, '_top_cal', None) or not self.date_entry._top_cal.winfo_ismapped():
                self.date_entry.drop_down()
            return "break"

        self.date_entry.bind("<Button-1>", _open_calendar_popup)
        self.date_entry.bind("<<DateEntrySelected>>", self._on_date_changed)

        # 3. Số tiền (VND)
        tk.Label(
            card_frame,
            text="💵  Số tiền (VND):",
            font=("Segoe UI", 9, "bold"),
            fg="#334155",
            bg="#FFFFFF"
        ).pack(anchor=tk.W, pady=(0, 3))

        self.amount_entry = tk.Entry(
            card_frame,
            font=("Segoe UI", 10),
            textvariable=self.amount_var,
            relief="solid",
            bd=1,
            highlightthickness=1,
            highlightcolor="#2563EB",
            highlightbackground="#CBD5E1"
        )
        self.amount_entry.pack(fill=tk.X, pady=(0, 10), ipady=4)

        # 4. Ghi chú thêm
        tk.Label(
            card_frame,
            text="📝  Ghi chú thêm:",
            font=("Segoe UI", 9, "bold"),
            fg="#334155",
            bg="#FFFFFF"
        ).pack(anchor=tk.W, pady=(0, 3))

        self.note_entry = tk.Entry(
            card_frame,
            font=("Segoe UI", 10),
            relief="solid",
            bd=1,
            highlightthickness=1,
            highlightcolor="#2563EB",
            highlightbackground="#CBD5E1"
        )
        self.note_entry.pack(fill=tk.X, pady=(0, 2), ipady=4)

        # C. KHU VỰC CÁC NÚT THAO TÁC (ACTION BUTTONS BO TRÒN)
        btn_container = tk.Frame(main_container, bg="#F8FAFC")
        btn_container.pack(fill=tk.X, pady=(2, 6))

        # Nút chính (Primary Action) - Bo góc 10px
        self.add_btn_norm, self.add_btn_hov = make_rounded_button_images(
            width=358, height=40, radius=10,
            fill_normal="#2563EB", fill_hover="#1D4ED8"
        )
        self.add_btn = tk.Button(
            btn_container,
            text="➕  Xác nhận Thêm & Đồng bộ",
            image=self.add_btn_norm,
            compound="center",
            font=("Segoe UI", 10, "bold"),
            fg="#FFFFFF",
            bg="#F8FAFC",
            activebackground="#F8FAFC",
            activeforeground="#FFFFFF",
            relief="flat",
            bd=0,
            highlightthickness=0,
            cursor="hand2",
            command=self.add_record
        )
        self.add_btn.pack(fill=tk.X, pady=(0, 8))
        self.add_btn.bind("<Enter>", lambda e: self.add_btn.config(image=self.add_btn_hov))
        self.add_btn.bind("<Leave>", lambda e: self.add_btn.config(image=self.add_btn_norm))

        # Hàng 2 nút phụ (Secondary Actions) - Bo góc 8px
        sec_btn_frame = tk.Frame(btn_container, bg="#F8FAFC")
        sec_btn_frame.pack(fill=tk.X)
        sec_btn_frame.columnconfigure(0, weight=1)
        sec_btn_frame.columnconfigure(1, weight=1)

        self.sync_btn_norm, self.sync_btn_hov = make_rounded_button_images(
            width=175, height=36, radius=8,
            fill_normal="#F1F5F9", fill_hover="#E2E8F0",
            outline_normal="#CBD5E1", outline_hover="#94A3B8"
        )
        self.sync_btn = tk.Button(
            sec_btn_frame,
            text="🔄  Tải dữ liệu mới",
            image=self.sync_btn_norm,
            compound="center",
            font=("Segoe UI", 9, "bold"),
            fg="#334155",
            bg="#F8FAFC",
            activebackground="#F8FAFC",
            activeforeground="#1E293B",
            relief="flat",
            bd=0,
            highlightthickness=0,
            cursor="hand2",
            command=self.manual_sync
        )
        self.sync_btn.grid(row=0, column=0, sticky="ew", padx=(0, 4))
        self.sync_btn.bind("<Enter>", lambda e: self.sync_btn.config(image=self.sync_btn_hov))
        self.sync_btn.bind("<Leave>", lambda e: self.sync_btn.config(image=self.sync_btn_norm))

        self.open_btn_norm, self.open_btn_hov = make_rounded_button_images(
            width=175, height=36, radius=8,
            fill_normal="#ECFDF5", fill_hover="#D1FAE5",
            outline_normal="#A7F3D0", outline_hover="#6EE7B7"
        )
        self.open_drive_btn = tk.Button(
            sec_btn_frame,
            text="📊  Mở Google Sheet",
            image=self.open_btn_norm,
            compound="center",
            font=("Segoe UI", 9, "bold"),
            fg="#065F46",
            bg="#F8FAFC",
            activebackground="#F8FAFC",
            activeforeground="#064E3B",
            relief="flat",
            bd=0,
            highlightthickness=0,
            cursor="hand2",
            command=self.open_file_in_browser
        )
        self.open_drive_btn.grid(row=0, column=1, sticky="ew", padx=(4, 0))
        self.open_drive_btn.bind("<Enter>", lambda e: self.open_drive_btn.config(image=self.open_btn_hov))
        self.open_drive_btn.bind("<Leave>", lambda e: self.open_drive_btn.config(image=self.open_btn_norm))

        # D. DÒNG THÔNG BÁO TRẠNG THÁI (STATUS FEEDBACK)
        self.status_msg = tk.Label(
            main_container,
            text="✨ Sẵn sàng nhập chi tiêu mới",
            font=("Segoe UI", 8),
            fg="#64748B",
            bg="#F8FAFC"
        )
        self.status_msg.pack(pady=(6, 2))

        # E. DÒNG CHỮ KÝ THIẾT KẾ & LIÊN KẾT WEB APP (FOOTER)
        self.designer_label = tk.Label(
            main_container,
            text="Designed & Developed by Trinh Ngoc Anh",
            font=("Segoe UI", 8, "italic"),
            fg="#94A3B8",
            bg="#F8FAFC"
        )
        self.designer_label.pack(side=tk.BOTTOM, pady=(2, 0))

        self.web_link_btn = tk.Label(
            main_container,
            text="📱 Web nhập liệu (Mobile): tna9899.github.io/Quanlychitieu",
            font=("Segoe UI", 8, "underline"),
            fg="#2563EB",
            bg="#F8FAFC",
            cursor="hand2"
        )
        self.web_link_btn.pack(side=tk.BOTTOM, pady=(4, 0))
        self.web_link_btn.bind("<Button-1>", lambda e: webbrowser.open("https://tna9899.github.io/Quanlychitieu/"))

        # Gán phím Enter để xác nhận Thêm & Đồng bộ (hỗ trợ cả phím Enter chính và Numpad Enter)
        self.root.bind("<Return>", self._handle_enter_key)
        self.root.bind("<KP_Enter>", self._handle_enter_key)

        # Mặc định đặt con trỏ vào ô nhập Số tiền để người dùng có thể nhập ngay
        self.amount_entry.focus_set()

        # Khởi chạy kết nối và tải dữ liệu trong luồng ngầm để giao diện hiển thị ngay lập tức
        self.root.after(50, self._start_initial_sync)

    def _on_date_changed(self, event=None):
        """Tự động đổi danh mục Điện nước -> Bố Mẹ khi chọn ngày từ tháng 10 trở đi"""
        try:
            sel_date = self.date_entry.get_date()
            cats = get_categories_for_month(sel_date.month, sel_date.year)
            cur = self.category_cb.get()
            self.category_cb['values'] = cats
            self.fixed_categories = cats
            if cur == "Điện nước" and "Bố Mẹ" in cats:
                self.category_cb.set("Bố Mẹ")
            elif cur == "Bố Mẹ" and "Điện nước" in cats:
                self.category_cb.set("Điện nước")
            elif cur not in cats:
                self.category_cb.set(cats[0])
        except Exception:
            pass

    def _bind_hover(self, widget, normal_bg, hover_bg):
        """Tạo hiệu ứng hover đổi màu nền cho nút bấm"""
        widget.bind("<Enter>", lambda e: widget.config(bg=hover_bg))
        widget.bind("<Leave>", lambda e: widget.config(bg=normal_bg))

    def _handle_enter_key(self, event=None):
        """Xử lý khi người dùng nhấn phím Enter để xác nhận Thêm & Đồng bộ"""
        # Nếu đang mở popup lịch chọn ngày thì để phím Enter cho lịch chọn ngày xử lý
        if hasattr(self, 'date_entry') and getattr(self.date_entry, '_top_cal', None):
            try:
                if self.date_entry._top_cal.winfo_exists() and self.date_entry._top_cal.winfo_ismapped():
                    return
            except Exception:
                pass

        # Nếu người dùng đang focus ở nút tải dữ liệu hoặc mở sheet, thực thi nút đó
        if event and event.widget:
            if event.widget == getattr(self, 'sync_btn', None):
                self.manual_sync()
                return
            if event.widget == getattr(self, 'open_drive_btn', None):
                self.open_file_in_browser()
                return

        self.add_record(event)

    def set_status(self, text, color="#64748B"):
        """Cập nhật thông điệp trạng thái cho người dùng (hỗ trợ gọi từ luồng phụ)"""
        def _update():
            if hasattr(self, 'status_msg') and self.status_msg:
                self.status_msg.config(text=text, fg=color)
        if threading.current_thread() is threading.main_thread():
            _update()
        else:
            self.root.after(0, _update)

    def _update_badge(self, text, fg, bg):
        """Cập nhật huy hiệu trạng thái an toàn đa luồng"""
        def _update():
            if hasattr(self, 'status_badge') and self.status_badge:
                self.status_badge.config(text=text, fg=fg, bg=bg)
        if threading.current_thread() is threading.main_thread():
            _update()
        else:
            self.root.after(0, _update)

    def _update_account_label(self, text):
        """Cập nhật email tài khoản an toàn đa luồng"""
        def _update():
            if hasattr(self, 'account_label') and self.account_label:
                self.account_label.config(text=text)
        if threading.current_thread() is threading.main_thread():
            _update()
        else:
            self.root.after(0, _update)

    def _start_initial_sync(self):
        """Khởi động luồng kết nối và nạp dữ liệu ban đầu không làm đơ giao diện"""
        self._is_syncing = True
        self.set_status("⏳ Đang kết nối Google Drive...", "#2563EB")
        self._update_badge("● Đang nạp...", "#F59E0B", "#FEF3C7")

        def _worker():
            self.reconnect_drive_with_existing_token()
            if self.drive_service:
                self.sync_from_drive()
            self.root.after(0, self._on_initial_sync_done)

        threading.Thread(target=_worker, daemon=True).start()

    def _on_initial_sync_done(self):
        """Hoàn tất nạp dữ liệu ban đầu"""
        self._is_syncing = False
        if self.drive_service:
            self._update_badge("● Đã kết nối", "#10B981", "#ECFDF5")
            self.set_status("✨ Sẵn sàng nhập chi tiêu mới", "#059669")
        else:
            self._update_badge("● Chưa kết nối", "#EF4444", "#FEF2F2")
            self.set_status("⚠️ Chưa thể kết nối Google Drive", "#EF4444")

    # ================= KHU VỰC XỬ LÝ GOOGLE DRIVE API =================
    def reconnect_drive_with_existing_token(self):
        # 1. Tìm credentials đã lưu: ưu tiên AppData, tiếp theo file nhúng, sau cùng là dữ liệu nhúng sẵn trong code
        creds = None
        app_data_token = os.path.join(get_app_data_dir(), 'token.json')
        embedded_token = get_resource_path('token.json')

        if os.path.exists(app_data_token):
            try:
                creds = Credentials.from_authorized_user_file(app_data_token, SCOPES)
            except Exception:
                pass
        
        if not creds and os.path.exists(embedded_token):
            try:
                creds = Credentials.from_authorized_user_file(embedded_token, SCOPES)
            except Exception:
                pass

        if not creds and DEFAULT_TOKEN_INFO.get("refresh_token"):
            try:
                creds = Credentials.from_authorized_user_info(DEFAULT_TOKEN_INFO, SCOPES)
            except Exception:
                pass

        # 2. Nếu token hết hạn hoặc chưa có, tự làm mới (refresh)
        if not creds or not creds.valid:
            if creds and creds.expired and creds.refresh_token:
                try:
                    creds.refresh(Request())
                    # Ghi đè lại token vào AppData để lần sau dùng tiếp
                    try:
                        with open(app_data_token, 'w', encoding='utf-8') as f:
                            f.write(creds.to_json())
                    except Exception:
                        pass
                except Exception:
                    creds = None
            
            if not creds:
                credentials_path = get_resource_path('credentials.json')
                if not os.path.exists(credentials_path):
                    # Sử dụng cấu hình nhúng sẵn nếu không tìm thấy file bên ngoài
                    credentials_path = os.path.join(get_app_data_dir(), 'credentials.json')
                    try:
                        with open(credentials_path, 'w', encoding='utf-8') as f:
                            json.dump(DEFAULT_CREDENTIALS_INFO, f)
                    except Exception:
                        pass

                try:
                    flow = InstalledAppFlow.from_client_secrets_file(credentials_path, SCOPES)
                    creds = flow.run_local_server(port=0)
                    try:
                        with open(app_data_token, 'w', encoding='utf-8') as f:
                            f.write(creds.to_json())
                    except Exception:
                        pass
                except Exception as e:
                    self._update_badge("● Lỗi xác thực", "#EF4444", "#FEF2F2")
                    self.root.after(0, lambda err=e: messagebox.showerror("Lỗi xác thực", f"Không thể kích hoạt trình duyệt đăng nhập: {err}", parent=self.root))
                    return

        # 3. Kết nối dịch vụ Google Drive
        try:
            self.drive_service = build('drive', 'v3', credentials=creds)
            self._update_badge("● Đã kết nối", "#10B981", "#ECFDF5")
            self.set_status("✨ Sẵn sàng nhập chi tiêu mới", "#059669")
            
            # Lấy thông tin tài khoản Google Drive đã đồng bộ
            try:
                about = self.drive_service.about().get(fields='user(displayName, emailAddress)').execute()
                u = about.get('user', {})
                email = u.get('emailAddress') or u.get('displayName', '')
                if email:
                    self._update_account_label(f"({email})")
            except Exception:
                pass
        except Exception as e:
            self._update_badge("● Lỗi kết nối", "#EF4444", "#FEF2F2")
            self.root.after(0, lambda err=e: messagebox.showerror("Lỗi kết nối API", f"Không thể kết nối Drive: {err}", parent=self.root))

    def _get_cached_file_id(self, year=None):
        """Đọc ID file Google Sheet theo năm từ AppData"""
        if year is None:
            year = getattr(self, 'current_year', datetime.now().year)
        year_str = str(year)
        
        # 1. Thử đọc từ drive_files_by_year.json
        cache_json_path = os.path.join(get_app_data_dir(), 'drive_files_by_year.json')
        if os.path.exists(cache_json_path):
            try:
                with open(cache_json_path, 'r', encoding='utf-8') as f:
                    mapping = json.load(f)
                    if isinstance(mapping, dict) and year_str in mapping and mapping[year_str]:
                        return mapping[year_str]
            except Exception:
                pass

        # 2. Backward compatibility: nếu là năm 2026 và có file drive_file_id.txt cũ hoặc DEFAULT_DRIVE_FILE_ID
        if year == 2026:
            cache_path = os.path.join(get_app_data_dir(), 'drive_file_id.txt')
            if os.path.exists(cache_path):
                try:
                    with open(cache_path, 'r', encoding='utf-8') as f:
                        fid = f.read().strip()
                        if fid:
                            return fid
                except Exception:
                    pass
            return DEFAULT_DRIVE_FILE_ID
            
        return None

    def _save_cached_file_id(self, file_id, year=None):
        """Lưu ID file Google Sheet theo năm vào AppData"""
        if not file_id:
            return
        if year is None:
            year = getattr(self, 'current_year', datetime.now().year)
        year_str = str(year)
        try:
            cache_json_path = os.path.join(get_app_data_dir(), 'drive_files_by_year.json')
            mapping = {}
            if os.path.exists(cache_json_path):
                try:
                    with open(cache_json_path, 'r', encoding='utf-8') as f:
                        mapping = json.load(f)
                except Exception:
                    mapping = {}
            mapping[year_str] = file_id
            with open(cache_json_path, 'w', encoding='utf-8') as f:
                json.dump(mapping, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def _clear_cached_file_id(self, year=None):
        """Xóa ID file trong cache khi bị lỗi hoặc không tìm thấy"""
        if year is None:
            year = getattr(self, 'current_year', datetime.now().year)
        year_str = str(year)
        try:
            cache_json_path = os.path.join(get_app_data_dir(), 'drive_files_by_year.json')
            if os.path.exists(cache_json_path):
                with open(cache_json_path, 'r', encoding='utf-8') as f:
                    mapping = json.load(f)
                if year_str in mapping:
                    del mapping[year_str]
                    with open(cache_json_path, 'w', encoding='utf-8') as f:
                        json.dump(mapping, f, ensure_ascii=False, indent=2)
        except Exception:
            pass
        if year == 2026:
            try:
                cache_path = os.path.join(get_app_data_dir(), 'drive_file_id.txt')
                if os.path.exists(cache_path):
                    os.remove(cache_path)
            except Exception:
                pass

    def find_or_create_file_on_drive(self, year=None):
        """Tìm hoặc tạo mới file trên Google Drive theo từng năm (chỉ tạo file mới khi là năm mới)"""
        if not self.drive_service:
            return None
        if year is None:
            year = getattr(self, 'current_year', datetime.now().year)

        target_name = get_drive_file_name_for_year(year)

        # 1. Thử lấy ID file từ cache local
        cached_id = self._get_cached_file_id(year)
        if cached_id:
            try:
                f_meta = self.drive_service.files().get(
                    fileId=cached_id, 
                    fields="id, name, trashed"
                ).execute()
                if f_meta and not f_meta.get('trashed', False):
                    # Nếu tên file trên Drive chưa có năm (ví dụ file cũ "Danh sách chi tiêu"), tự động đổi tên thành "Danh sách chi tiêu {year}"
                    if f_meta.get('name') != target_name and year == 2026:
                        try:
                            self.drive_service.files().update(
                                fileId=cached_id, 
                                body={'name': target_name}
                            ).execute()
                        except Exception:
                            pass
                    self.file_id_on_drive = cached_id
                    self._save_cached_file_id(cached_id, year)
                    return self.file_id_on_drive
            except Exception as e:
                print(f"File cache ID không hợp lệ cho năm {year} ({cached_id}): {e}")
                self._clear_cached_file_id(year)

        # 2. Tìm kiếm trên Google Drive theo tên file của năm đó
        try:
            results = self.drive_service.files().list(
                q=f"name='{target_name}' and mimeType='application/vnd.google-apps.spreadsheet' and trashed=false",
                fields="files(id, name)"
            ).execute()
            items = results.get('files', [])
            
            if items:
                self.file_id_on_drive = items[0]['id']
                self._save_cached_file_id(self.file_id_on_drive, year)
                return self.file_id_on_drive

            # 3. Nếu là năm 2026, thử tìm file cũ không có đuôi năm "Danh sách chi tiêu"
            if year == 2026:
                old_results = self.drive_service.files().list(
                    q=f"name='{DRIVE_FILE_BASE_NAME}' and mimeType='application/vnd.google-apps.spreadsheet' and trashed=false",
                    fields="files(id, name)"
                ).execute()
                old_items = old_results.get('files', [])
                if old_items:
                    self.file_id_on_drive = old_items[0]['id']
                    try:
                        self.drive_service.files().update(
                            fileId=self.file_id_on_drive, 
                            body={'name': target_name}
                        ).execute()
                    except Exception:
                        pass
                    self._save_cached_file_id(self.file_id_on_drive, year)
                    return self.file_id_on_drive

            # 4. CHỈ KHI LÀ NĂM MỚI (chưa có file trên Drive), TẠO MỚI FILE CHO NĂM ĐÓ
            file_metadata = {
                'name': target_name,
                'mimeType': 'application/vnd.google-apps.spreadsheet'
            }
            empty_df = pd.DataFrame(columns=["Ngày", "Danh mục", "Số tiền", "Ghi chú"])
            towrite = io.BytesIO()
            month_to_init = datetime.now().strftime('%m') if year == datetime.now().year else '01'
            current_month_sheet = f"Tháng {month_to_init}-{year}"
            with pd.ExcelWriter(towrite, engine='openpyxl') as writer:
                empty_df.to_excel(writer, sheet_name=current_month_sheet, index=False)
            
            # Sửa định dạng upload: stream là openxmlformats, Drive tự convert sang Google Spreadsheet
            media = MediaInMemoryUpload(
                towrite.getvalue(), 
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            file = self.drive_service.files().create(
                body=file_metadata, 
                media_body=media, 
                fields='id'
            ).execute()
            self.file_id_on_drive = file.get('id')
            self._save_cached_file_id(self.file_id_on_drive, year)
            return self.file_id_on_drive
        except Exception as e:
            print(f"Lỗi tìm/tạo file đám mây cho năm {year}: {e}")
            return None

    def sync_from_drive(self, year=None):
        """Tải dữ liệu mới nhất từ Google Drive về máy cho năm được chọn"""
        if not self.drive_service:
            return
        if year is None:
            year = getattr(self, 'current_year', datetime.now().year)

        fid = self.find_or_create_file_on_drive(year)
        if not fid:
            return
        self.file_id_on_drive = fid
        try:
            request = self.drive_service.files().export_media(
                fileId=self.file_id_on_drive,
                mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            fh = io.BytesIO()
            downloader = MediaIoBaseDownload(fh, request)
            done = False
            while done is False:
                status, done = downloader.next_chunk()
            
            fh.seek(0)
            self.load_data_from_stream(fh)
        except Exception as e:
            print(f"Lỗi đồng bộ tải xuống năm {year}: {e}")
            if "404" in str(e) or "notFound" in str(e):
                self.file_id_on_drive = None
                self._clear_cached_file_id(year)
                self.find_or_create_file_on_drive(year)

    def upload_to_drive(self, target_year=None):
        if not self.drive_service:
            return
        if target_year is None:
            target_year = getattr(self, 'current_year', datetime.now().year)

        fid = self.find_or_create_file_on_drive(target_year)
        if not fid:
            raise RuntimeError(f"Không tìm thấy hoặc không thể tạo file trên Google Drive cho năm {target_year}")
        self.file_id_on_drive = fid

        try:
            if self.df.empty:
                return

            self.df["Ngày"] = pd.to_datetime(self.df["Ngày"])
            
            # Lọc dữ liệu thuộc đúng năm target_year để lưu vào file của năm đó
            df_year = self.df[self.df["Ngày"].dt.year == target_year].copy()
            if df_year.empty:
                return

            df_year = df_year.sort_values(by="Ngày").reset_index(drop=True)
            
            # Đồng bộ định dạng tên tháng (ví dụ: Tháng 06-2026, Tháng 07-2026)
            df_year["Month_Group"] = df_year["Ngày"].dt.strftime("Tháng %m-%Y")
            
            current_month_str = datetime.now().strftime(f"Tháng %m-{target_year}")
            unique_months = list(df_year["Month_Group"].unique())
            
            # Sắp xếp để đưa tab tháng hiện tại lên đầu tiên
            if current_month_str in unique_months:
                unique_months.remove(current_month_str)
                unique_months.insert(0, current_month_str)

            wb = openpyxl.Workbook()
            default_sheet = wb.active
            wb.remove(default_sheet)

            black_75_side = Side(border_style="thin", color="404040")
            data_borders = Border(left=black_75_side, right=black_75_side, top=black_75_side, bottom=black_75_side)
            
            center_align = Alignment(horizontal="center", vertical="center")
            left_align = Alignment(horizontal="left", vertical="center")
            right_align = Alignment(horizontal="right", vertical="center") 
            
            header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            data_font = Font(name="Arial", size=10)
            total_font = Font(name="Arial", size=10, bold=True, color="D32F2F")

            header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid") 
            summary_header_fill = PatternFill(start_color="43A047", end_color="43A047", fill_type="solid") 

            # VÒNG LẶP GENERATE LẠI TOÀN BỘ CÁC SHEET THÁNG CÓ DỮ LIỆU CỦA NĂM
            for month_name in unique_months:
                group = df_year[df_year["Month_Group"] == month_name]
                ws = wb.create_sheet(title=month_name)
                ws.views.sheetView[0].showGridLines = False 

                # 1. Tiêu đề
                headers = ["Ngày", "Danh mục", "Số tiền", "Ghi chú"]
                ws.append(headers)
                for col_idx in range(1, 5):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align
                    cell.border = data_borders 
                ws.row_dimensions[1].height = 25

                # 2. Ghi dữ liệu chi tiết
                for _, row in group.iterrows():
                    date_str = row["Ngày"].strftime("%Y-%m-%d")
                    ws.append([date_str, row["Danh mục"], row["Số tiền"], row["Ghi chú"]])
                    curr_row = ws.max_row
                    ws.row_dimensions[curr_row].height = 20
                    
                    for col_idx in range(1, 5):
                        c = ws.cell(row=curr_row, column=col_idx)
                        c.font = data_font
                        c.border = data_borders 
                        if col_idx in [1, 2]:
                            c.alignment = center_align
                        else:
                            c.alignment = left_align
                        
                        if col_idx == 3:
                            c.alignment = right_align
                            c.number_format = '#,##0'

                # 3. Khối tổng hợp rút gọn (F và G)
                try:
                    month_num = int(month_name.replace("Tháng ", "").split("-")[0])
                except Exception:
                    month_num = 10
                current_month_cats = get_categories_for_month(month_num, target_year)

                ws["F1"] = "Danh mục tổng"
                ws["G1"] = "Tổng chi tiêu"
                for col_letter in ["F1", "G1"]:
                    ws[col_letter].font = header_font
                    ws[col_letter].fill = summary_header_fill
                    ws[col_letter].alignment = center_align
                    ws[col_letter].border = data_borders 

                for idx, cat_name in enumerate(current_month_cats):
                    row_num = 2 + idx
                    ws[f"F{row_num}"] = cat_name
                    ws[f"G{row_num}"] = f"=SUMIF(B:B, F{row_num}, C:C)"
                    
                    for col_letter in [f"F{row_num}", f"G{row_num}"]:
                        ws[col_letter].font = data_font
                        ws[col_letter].border = data_borders
                        if col_letter.startswith("F"):
                            ws[col_letter].alignment = center_align
                        else:
                            ws[col_letter].alignment = right_align
                            ws[col_letter].number_format = '#,##0'
                    ws.row_dimensions[row_num].height = 20

                # 4. Ô Tổng cộng tháng
                last_cat_row = 1 + len(current_month_cats)
                total_row = last_cat_row + 1

                ws[f"F{total_row}"] = "Tổng trong tháng"
                ws[f"G{total_row}"] = f"=SUM(G2:G{last_cat_row})" 
                
                ws[f"F{total_row}"].font = total_font
                ws[f"F{total_row}"].border = data_borders 
                ws[f"F{total_row}"].alignment = center_align

                ws[f"G{total_row}"].font = total_font
                ws[f"G{total_row}"].border = data_borders 
                ws[f"G{total_row}"].alignment = right_align 
                ws[f"G{total_row}"].number_format = '#,##0'
                ws.row_dimensions[total_row].height = 22

                # 5. Biểu đồ tròn
                pie = PieChart()
                pie.title = f"Tỷ lệ chi tiêu - {month_name}"
                data_ref = Reference(ws, min_col=7, min_row=1, max_row=last_cat_row)
                labels_ref = Reference(ws, min_col=6, min_row=2, max_row=last_cat_row)
                
                pie.add_data(data_ref, titles_from_data=True)
                pie.set_categories(labels_ref)
                pie.style = 10  
                pie.width = 11
                pie.height = 8
                ws.add_chart(pie, "I1")

                # Cấu hình độ rộng cột tiêu chuẩn tối ưu
                ws.column_dimensions["A"].width = 14
                ws.column_dimensions["B"].width = 15
                ws.column_dimensions["C"].width = 16
                ws.column_dimensions["D"].width = 25
                ws.column_dimensions["F"].width = 18
                ws.column_dimensions["G"].width = 18

            towrite = io.BytesIO()
            wb.save(towrite)
            wb.close()
            
            # resumable=False tải lên trực tiếp trong 1 request nhanh hơn nhiều lần
            media = MediaInMemoryUpload(
                towrite.getvalue(), 
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 
                resumable=False
            )
            self.drive_service.files().update(fileId=self.file_id_on_drive, media_body=media).execute()
        except Exception as e:
            raise RuntimeError(f"Không thể upload đồng bộ lên Google Drive: {e}")

    def load_data_from_stream(self, stream):
        required_cols = ["Ngày", "Danh mục", "Số tiền", "Ghi chú"]
        all_sheets_data = []

        try:
            with pd.ExcelFile(stream) as excel_file:
                for sheet_name in excel_file.sheet_names:
                    # ĐỌC TOÀN BỘ CÁC SHEET: Gom sạch dữ liệu từ mọi sheet chứa chữ "tháng" (Tháng 06, Tháng 07,...)
                    if "tháng" not in sheet_name.lower():
                        continue
                    
                    df_sheet = pd.read_excel(excel_file, sheet_name=sheet_name, usecols="A:D")
                    df_sheet = df_sheet.dropna(subset=["Ngày", "Danh mục", "Số tiền"], how="all")
                    
                    for col in required_cols:
                        if col not in df_sheet.columns:
                            df_sheet[col] = 0 if col == "Số tiền" else ""
                    
                    df_sheet = df_sheet[required_cols]
                    all_sheets_data.append(df_sheet)
            
            if all_sheets_data:
                self.df = pd.concat(all_sheets_data, ignore_index=True)
                self.df["Ngày"] = pd.to_datetime(self.df["Ngày"], errors='coerce')
                self.df["Ngày"] = self.df["Ngày"].fillna(pd.to_datetime(datetime.now().date()))
                self.df["Số tiền"] = pd.to_numeric(self.df["Số tiền"], errors='coerce').fillna(0)
                self.df = self.df[self.df["Danh mục"].notna() & (self.df["Số tiền"] > 0)]
            else:
                self.df = pd.DataFrame(columns=required_cols)
        except Exception as e:
            print(f"Lỗi phân tích dữ liệu luồng: {e}")
            self.df = pd.DataFrame(columns=required_cols)

    def manual_sync(self):
        if not self.drive_service:
            messagebox.showwarning("Cảnh báo", "Chưa kết nối được Google Drive!", parent=self.root)
            return
        if self._is_syncing:
            messagebox.showwarning("Đang xử lý", "Hệ thống đang đồng bộ dữ liệu, vui lòng đợi trong giây lát!", parent=self.root)
            return

        self._is_syncing = True
        self.set_status("⏳ Đang tải dữ liệu mới từ Google Drive...", "#2563EB")
        self.sync_btn.config(state="disabled", cursor="watch")

        year = getattr(self, 'current_year', datetime.now().year)

        def _worker():
            err_msg = None
            try:
                self.sync_from_drive(year)
            except Exception as e:
                err_msg = str(e)
            self.root.after(0, lambda: self._on_manual_sync_done(err_msg))

        threading.Thread(target=_worker, daemon=True).start()

    def _on_manual_sync_done(self, err_msg=None):
        self._is_syncing = False
        self.sync_btn.config(state="normal", cursor="hand2")
        year = getattr(self, 'current_year', datetime.now().year)
        if err_msg:
            self.set_status("❌ Lỗi tải dữ liệu!", "#EF4444")
            messagebox.showerror("Lỗi đồng bộ", f"Không thể tải từ Google Drive: {err_msg}", parent=self.root)
        else:
            self.set_status(f"✅ Đã cập nhật dữ liệu năm {year} từ Google Sheets!", "#059669")
            messagebox.showinfo("Đồng bộ thành công", f"Đã cập nhật đầy đủ dữ liệu tất cả các tháng của năm {year} từ Google Sheets!", parent=self.root)

    def open_file_in_browser(self):
        if not self.drive_service:
            messagebox.showwarning("Cảnh báo", "Chưa kết nối dịch vụ Google Drive!", parent=self.root)
            return
        year = getattr(self, 'current_year', datetime.now().year)
        fid = self.find_or_create_file_on_drive(year)
        if fid:
            url = f"https://docs.google.com/spreadsheets/d/{fid}/edit"
            webbrowser.open(url)
        else:
            messagebox.showerror("Lỗi", f"Không tìm thấy ID file năm {year} trên Google Drive để mở!", parent=self.root)

    def format_amount_live(self, *args):
        text = self.amount_var.get()
        cleaned = "".join([c for c in text if c.isdigit()])
        if cleaned == "":
            if text != "": self.amount_var.set("")
            return
        formatted = f"{int(cleaned):,}"
        if text != formatted:
            self.amount_var.set(formatted)
            self.amount_entry.icursor(tk.END)

    def add_record(self, event=None):
        if self._is_syncing:
            messagebox.showwarning("Đang xử lý", "Hệ thống đang đồng bộ dữ liệu, vui lòng đợi trong giây lát!", parent=self.root)
            return

        category = self.category_cb.get().strip()
        selected_date = self.date_entry.get_date()
        amount_str = self.amount_entry.get().strip().replace(",", "")
        note = self.note_entry.get().strip()

        if not category:
            messagebox.showwarning("Cảnh báo", "Vui lòng chọn Danh mục chi tiêu!", parent=self.root)
            return
        try:
            amount = float(amount_str)
            if amount <= 0: raise ValueError
        except ValueError:
            messagebox.showerror("Lỗi", "Số tiền phải là số dương hợp lệ!", parent=self.root)
            return

        rec_year = selected_date.year

        # Nếu năm của giao dịch khác năm đang nạp, đồng bộ dữ liệu của năm đó trước
        if rec_year != getattr(self, 'current_year', datetime.now().year):
            self.current_year = rec_year
            try:
                self.sync_from_drive(rec_year)
            except Exception:
                pass

        new_row = pd.DataFrame([{
            "Ngày": pd.to_datetime(selected_date), "Danh mục": category, "Số tiền": amount, "Ghi chú": note
        }])
        
        # Thêm ngay vào DataFrame cục bộ và xóa trường nhập để người dùng sẵn sàng nhập tiếp
        self.df = pd.concat([self.df, new_row], ignore_index=True)
        self.amount_var.set("")
        self.note_entry.delete(0, tk.END)

        # Chuyển trạng thái giao diện sang "Đang lưu..." (xử lý ngầm không làm đơ ứng dụng)
        self._is_syncing = True
        self.set_status("⏳ Đang lưu & đồng bộ lên Google Drive...", "#2563EB")
        self.add_btn.config(text="⏳  Đang lưu dữ liệu...", state="disabled", cursor="watch")

        def _worker():
            err_msg = None
            try:
                self.upload_to_drive(rec_year)
            except Exception as e:
                err_msg = str(e)
            self.root.after(0, lambda: self._on_add_done(err_msg))

        threading.Thread(target=_worker, daemon=True).start()

    def _on_add_done(self, err_msg=None):
        self._is_syncing = False
        self.add_btn.config(text="➕  Xác nhận Thêm & Đồng bộ", state="normal", cursor="hand2")
        if err_msg:
            self.set_status("❌ Lỗi lưu dữ liệu!", "#EF4444")
            messagebox.showerror("Lỗi Cloud", err_msg, parent=self.root)
        else:
            now_time = datetime.now().strftime("%H:%M:%S")
            self.set_status(f"✅ Đã lưu và đồng bộ lên Google Drive! ({now_time})", "#059669")
            messagebox.showinfo("Thành công", "Đã lưu dữ liệu lên Google Drive!", parent=self.root)
            self.amount_entry.focus_set()


if __name__ == "__main__":
    # Đăng ký App ID ngay trước khi khởi tạo root window (chỉ khi chạy python dev)
    if sys.platform.startswith('win'):
        if not getattr(sys, 'frozen', False):
            try:
                app_id = 'trinhngocanh.quanlychitieu.app.1.0'
                ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(app_id)
            except Exception:
                pass

    main_root = tk.Tk()
    app = ExpenditureApp(main_root)
    main_root.mainloop()